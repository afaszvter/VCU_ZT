from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "Module" / "说明文档" / "ZT_VehCtrl_24B_2_ICU_DBC信号对应表.md"
OUTPUT_XLSX = ROOT / "Module" / "说明文档" / "ZT_VehCtrl_24B_2_ICU_DBC信号对应表_更新版.xlsx"


def clean_markdown(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    return re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text).strip()


def parse_report(path: Path) -> list[dict[str, object]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    message = ""
    can_id = ""
    records: list[dict[str, object]] = []
    heading_re = re.compile(r"^###\s+(VCU2ICUInfor\d+)（(0x[0-9A-Fa-f]+)）$")

    for line in lines:
        heading = heading_re.match(line.strip())
        if heading:
            message, can_id = heading.groups()
            continue

        if not message or not re.match(r"^\|\s*\d+\s*\|", line):
            continue

        cells = [clean_markdown(cell) for cell in line.strip().strip("|").split("|")]
        if len(cells) != 5:
            raise ValueError(f"无法解析表格行（列数={len(cells)}）：{line}")

        local_order, signal, source, detail, status = cells
        if status == "空缺：固定0":
            source_out = ""
            condition = ""
            processing = "CAN Pack 前固定 0（模型未接入）"
        else:
            source_out = source
            # 只有真正的逻辑/阈值判断进入“判断条件”；转换与聚合放到备注列。
            if signal == "VCU_SOCLoFlt_b":
                condition = detail
                processing = ""
            elif signal == "VCU_BattTempHi_b":
                condition = detail
                processing = ""
            elif signal == "VCU_MCUTempHiFlt_b":
                condition = "max(8路MCU温度) >= MCUTempAlarmTh（60℃）"
                processing = "8 路温度取最大值"
            elif signal == "VCU_MtTempHiFlt_b":
                condition = "max(8路电机温度) >= MtrTempAlarmTh（50℃）"
                processing = "8 路温度取最大值"
            elif signal == "VCU_Pst_b":
                condition = detail
                processing = ""
            else:
                condition = ""
                processing = "" if detail in {"直接映射", "一致"} else detail

        records.append(
            {
                "global_order": len(records) + 1,
                "message": message,
                "can_id": can_id.upper().replace("0X", "0x"),
                "local_order": int(local_order),
                "signal": signal,
                "source": source_out,
                "condition": condition,
                "processing": processing,
                "status": status,
            }
        )

    if len(records) != 58:
        raise ValueError(f"期望 58 条信号，实际解析到 {len(records)} 条")
    return records


def add_main_sheet(wb: Workbook, records: list[dict[str, object]]) -> None:
    ws = wb.active
    ws.title = "信号对应表"
    ws.sheet_view.showGridLines = False

    navy = "17365D"
    blue = "1F4E78"
    light_blue = "D9EAF7"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="D9E1F2")
    medium_blue = Side(style="medium", color=blue)

    ws.merge_cells("A1:I1")
    ws["A1"] = "ZT_VehCtrl_24B_2 — ICU 与 DBC 信号对应表"
    ws["A1"].font = Font(name="微软雅黑", size=16, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:I2")
    ws["A2"] = "顺序严格遵循 DBC 的 BO_/SG_ 出现顺序；来源空白表示模型未接入，实际发送固定 0"
    ws["A2"].font = Font(name="微软雅黑", size=10, italic=True, color="404040")
    ws["A2"].fill = PatternFill("solid", fgColor="EAF2F8")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 24

    headers = [
        "全局序号",
        "报文",
        "CAN ID",
        "DBC序号",
        "ICU信号名",
        "模型实际来源（空白=未接入）",
        "判断条件",
        "信号处理/备注",
        "状态",
    ]
    header_row = 4
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, value)
        cell.font = Font(name="微软雅黑", size=10.5, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    ws.row_dimensions[header_row].height = 34

    status_fills = {
        "一致": PatternFill("solid", fgColor="E2F0D9"),
        "空缺：固定0": PatternFill("solid", fgColor="FCE4D6"),
        "错位": PatternFill("solid", fgColor="FFC7CE"),
        "需确认": PatternFill("solid", fgColor="FFF2CC"),
    }
    status_fonts = {
        "一致": "375623",
        "空缺：固定0": "9C0006",
        "错位": "9C0006",
        "需确认": "7F6000",
    }

    first_data_row = header_row + 1
    previous_message = None
    for row_idx, record in enumerate(records, start=first_data_row):
        values = [
            record["global_order"],
            record["message"],
            record["can_id"],
            record["local_order"],
            record["signal"],
            record["source"],
            record["condition"],
            record["processing"],
            record["status"],
        ]
        is_new_message = previous_message is not None and record["message"] != previous_message
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = Font(name="微软雅黑", size=10, color="000000")
            cell.alignment = Alignment(
                horizontal="center" if col_idx in {1, 2, 3, 4, 9} else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(
                left=thin_gray,
                right=thin_gray,
                top=medium_blue if is_new_message else thin_gray,
                bottom=thin_gray,
            )
        ws.cell(row_idx, 5).font = Font(name="Consolas", size=10, color="1F1F1F")
        ws.cell(row_idx, 9).fill = status_fills[str(record["status"])]
        ws.cell(row_idx, 9).font = Font(
            name="微软雅黑", size=10, bold=True, color=status_fonts[str(record["status"])]
        )
        ws.row_dimensions[row_idx].height = 36 if record["status"] in {"错位", "需确认"} else 30
        previous_message = record["message"]

    last_row = first_data_row + len(records) - 1
    table = Table(displayName="ICU_DBC_Signal_Map", ref=f"A{header_row}:I{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)

    # Reinforce status coloring when users append or edit rows in Excel.
    for status, fill in status_fills.items():
        ws.conditional_formatting.add(
            f"I{first_data_row}:I{last_row}",
            FormulaRule(formula=[f'$I{first_data_row}="{status}"'], fill=fill),
        )

    widths = {"A": 10, "B": 20, "C": 14, "D": 10, "E": 30, "F": 62, "G": 46, "H": 72, "I": 16}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:I{last_row}"
    ws.print_title_rows = f"1:{header_row}"
    ws.print_area = f"A1:I{last_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def add_notes_sheet(wb: Workbook, records: list[dict[str, object]]) -> None:
    ws = wb.create_sheet("说明")
    ws.sheet_view.showGridLines = False
    navy = "17365D"
    blue = "1F4E78"
    light_blue = "D9EAF7"
    white = "FFFFFF"
    border_side = Side(style="thin", color="D9E1F2")

    ws.merge_cells("A1:D1")
    ws["A1"] = "使用说明与核对结论"
    ws["A1"].font = Font(name="微软雅黑", size=15, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    status_counts = Counter(str(record["status"]) for record in records)
    content = [
        ("模型", r"Module\MainModel\ZT_VehCtrl_24B_2.slx"),
        ("分析对象", "ZT_VehCtrl_24B_2/VehCtrl/ICU 的 7 条 VCU→ICU 报文"),
        ("DBC顺序依据", r"Module\CAN_DBC\有效CAN信号\整车CANA_GBK._Valid.dbc"),
        ("原始分析报告", r"Module\说明文档\ZT_VehCtrl_24B_2_ICU_DBC信号对应表.md"),
        ("核对日期", "2026-09-29"),
        ("结果快照", f"一致 {status_counts['一致']}；空缺：固定0 {status_counts['空缺：固定0']}；错位 {status_counts['错位']}；需确认 {status_counts['需确认']}；合计 {len(records)}"),
        ("字段约定", "简单直连、类型转换、缩放、限幅和取最大值等不写入“判断条件”，统一放在“信号处理/备注”；只有逻辑或阈值判定写入“判断条件”。"),
        ("空缺约定", "DBC 有定义但模型固定发送 0 的信号，其“模型实际来源”和“判断条件”保持空白，状态标为“空缺：固定0”。"),
        ("修正确认", "VCU2ICUInfor5 与 VCU2ICUInfor6 的 8 个 CAN Pack 输入端口已重新核对，当前接线均与端口标签一致。"),
        ("优先问题 2", "VCU_WrkSt_enum 实际发送 EnergyManage.Ready；模型内另有同名故障计算信号但未接入 CAN Pack，需要策略确认。"),
        ("限制", "“一致”仅表示模型接线、CAN Pack 端口标签和 DBC 字段对应，不代表传感器标定或实车物理语义已验证。"),
    ]

    for row_idx, (label, value) in enumerate(content, start=3):
        ws.cell(row_idx, 1, label)
        ws.cell(row_idx, 2, value)
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
        for col_idx in range(1, 5):
            cell = ws.cell(row_idx, col_idx)
            cell.font = Font(name="微软雅黑", size=10, bold=(col_idx == 1))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
            if col_idx == 1:
                cell.fill = PatternFill("solid", fgColor=light_blue)
        ws.row_dimensions[row_idx].height = 33 if row_idx >= 9 else 25

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 42
    ws.freeze_panes = "A3"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def verify_workbook(path: Path, records: list[dict[str, object]]) -> None:
    wb = load_workbook(path, data_only=False)
    if wb.sheetnames != ["信号对应表", "说明"]:
        raise AssertionError(f"工作表异常：{wb.sheetnames}")
    ws = wb["信号对应表"]
    actual = []
    for row in ws.iter_rows(min_row=5, max_row=62, values_only=True):
        actual.append(
            {
                "message": row[1],
                "can_id": row[2],
                "local_order": row[3],
                "signal": row[4],
                "status": row[8],
            }
        )
    expected = [
        {key: record[key] for key in ("message", "can_id", "local_order", "signal", "status")}
        for record in records
    ]
    if actual != expected:
        raise AssertionError("工作簿信号顺序或内容与源报告不一致")
    if ws.freeze_panes != "A5":
        raise AssertionError(f"冻结窗格异常：{ws.freeze_panes}")
    if ws.tables["ICU_DBC_Signal_Map"].ref != "A4:I62":
        raise AssertionError("Excel 表格范围异常")


def main() -> None:
    records = parse_report(SOURCE_MD)
    wb = Workbook()
    wb.properties.title = "ZT_VehCtrl_24B_2 ICU 与 DBC 信号对应表"
    wb.properties.subject = "7 条 VCU→ICU 报文的 58 个信号模型来源核对"
    wb.properties.creator = "Codex"
    wb.properties.description = "按 DBC 顺序整理；包含模型来源、判断条件、信号处理和核对状态。"
    add_main_sheet(wb, records)
    add_notes_sheet(wb, records)
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)
    verify_workbook(OUTPUT_XLSX, records)
    counts = Counter(str(record["status"]) for record in records)
    print(f"OUTPUT={OUTPUT_XLSX}")
    print(f"ROWS={len(records)}")
    print("STATUS=" + ", ".join(f"{key}:{counts[key]}" for key in ["一致", "空缺：固定0", "错位", "需确认"]))


if __name__ == "__main__":
    main()
