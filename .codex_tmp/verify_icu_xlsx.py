from __future__ import annotations

import re
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Module" / "说明文档" / "ZT_VehCtrl_24B_2_ICU_DBC信号对应表_更新版.xlsx"
DBC = ROOT / "Module" / "CAN_DBC" / "有效CAN信号" / "整车CANA_GBK._Valid.dbc"


def dbc_order() -> list[tuple[str, str, int, str]]:
    text = DBC.read_text(encoding="gbk", errors="replace")
    target = re.compile(r"^BO_\s+(\d+)\s+(VCU2ICUInfor[1-7]):")
    signal = re.compile(r"^\s+SG_\s+(\w+)\s*:")
    rows: list[tuple[str, str, int, str]] = []
    message = None
    can_id = None
    local_order = 0
    for line in text.splitlines():
        match = target.match(line)
        if match:
            raw_id, message = match.groups()
            can_id = f"0x{int(raw_id) & 0x1FFFFFFF:08X}"
            local_order = 0
            continue
        if line.startswith("BO_"):
            message = None
            can_id = None
            continue
        sig = signal.match(line)
        if message and can_id and sig:
            local_order += 1
            rows.append((message, can_id, local_order, sig.group(1)))
    return rows


def main() -> None:
    if not XLSX.exists():
        raise AssertionError(f"缺少输出文件：{XLSX}")
    if not zipfile.is_zipfile(XLSX):
        raise AssertionError("输出文件不是有效的 XLSX/ZIP 包")
    with zipfile.ZipFile(XLSX) as archive:
        required = {"[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml"}
        missing = required.difference(archive.namelist())
        if missing:
            raise AssertionError(f"XLSX 包缺少必要部件：{sorted(missing)}")

    wb = load_workbook(XLSX, data_only=False)
    if wb.sheetnames != ["信号对应表", "说明"]:
        raise AssertionError(f"工作表名称/顺序异常：{wb.sheetnames}")
    ws = wb["信号对应表"]
    data = list(ws.iter_rows(min_row=5, max_row=62, values_only=True))
    if len(data) != 58:
        raise AssertionError(f"信号行数不是 58：{len(data)}")

    actual_order = [(row[1], row[2], row[3], row[4]) for row in data]
    expected_order = dbc_order()
    if actual_order != expected_order:
        for index, (actual, expected) in enumerate(zip(actual_order, expected_order), start=1):
            if actual != expected:
                raise AssertionError(f"DBC 顺序第 {index} 项不一致：Excel={actual}，DBC={expected}")
        raise AssertionError(f"DBC/Excel 数量不同：Excel={len(actual_order)}，DBC={len(expected_order)}")

    counts = Counter(row[8] for row in data)
    expected_counts = Counter({"一致": 34, "空缺：固定0": 23, "错位": 0, "需确认": 1})
    if counts != expected_counts:
        raise AssertionError(f"状态计数异常：{counts}")

    for row in data:
        status = row[8]
        source = row[5]
        condition = row[6]
        if status == "空缺：固定0" and (source is not None or condition is not None):
            raise AssertionError(f"空缺信号未按要求留空：{row[4]}")
        if status != "空缺：固定0" and not source:
            raise AssertionError(f"已接入信号缺少来源：{row[4]}")

    condition_signals = {row[4] for row in data if row[6]}
    expected_condition_signals = {
        "VCU_Pst_b",
        "VCU_SOCLoFlt_b",
        "VCU_BattTempHi_b",
        "VCU_MCUTempHiFlt_b",
        "VCU_MtTempHiFlt_b",
    }
    if condition_signals != expected_condition_signals:
        raise AssertionError(f"判断条件列信号集合异常：{sorted(condition_signals)}")

    corrected_sources = {
        "VCU_DrvMtSpd_rpm": "TM：Motor_Speed；MCU2：DrvMotorSpd",
        "VCU_DrvMtTrq_Nm": "TM：Motor_Torque；MCU2：DrvMotorTrq",
        "VCU_DrvMCUBusVolt_V": "TM：DC_Voltage_A；MCU2：DrvMotorBusVolt",
        "VCU_DrvMCUBusCurr_A": "TM：DC_Current_A；MCU2：DrvMotorBusCur",
        "VCU_OilMtSpd_rpm": "ISG：Motor_Speed；MCU1：OilMotorSpd",
        "VCU_OilMtTrq_Nm": "ISG：Motor_Torque；MCU1：OilMotorTrq",
        "VCU_OilMCUBusVolt_V": "ISG：DC_Voltage_V；MCU1：OilMotorBusVolt",
        "VCU_OilMCUBusCurr_A": "ISG：DC_Current_A；MCU1：OilMotorBusCur",
    }
    corrected_rows = {row[4]: row for row in data if row[4] in corrected_sources}
    if set(corrected_rows) != set(corrected_sources):
        raise AssertionError("第 5、6 条报文的 8 个目标信号不完整")
    for signal_name, expected_source in corrected_sources.items():
        row = corrected_rows[signal_name]
        if row[5] != expected_source or row[7] != "两路取最大值" or row[8] != "一致":
            raise AssertionError(
                f"修正信号内容异常：{signal_name}, source={row[5]}, process={row[7]}, status={row[8]}"
            )

    formula_cells = [
        cell.coordinate
        for sheet in wb.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    if formula_cells:
        raise AssertionError(f"存在未经外部重算的公式：{formula_cells}")
    if ws.freeze_panes != "A5":
        raise AssertionError(f"冻结窗格异常：{ws.freeze_panes}")
    if ws.tables["ICU_DBC_Signal_Map"].ref != "A4:I62":
        raise AssertionError("Excel 表范围不是 A4:I62")
    if ws.auto_filter.ref != "A4:I62":
        raise AssertionError(f"筛选范围异常：{ws.auto_filter.ref}")

    print(f"PASS file={XLSX}")
    print(f"PASS xlsx_package_parts={len(zipfile.ZipFile(XLSX).namelist())}")
    print(f"PASS sheets={wb.sheetnames}")
    print(f"PASS dbc_order_rows={len(expected_order)} mismatch=0")
    print(f"PASS status_counts={dict(counts)}")
    print("PASS fixed0_blank_fields=23/23")
    print(f"PASS condition_signals={sorted(condition_signals)}")
    print("PASS ICU5_6_corrected_rows=8/8 status=一致")
    print("PASS formulas=0")
    print("PASS table=A4:I62 freeze=A5 filter=A4:I62")


if __name__ == "__main__":
    main()
